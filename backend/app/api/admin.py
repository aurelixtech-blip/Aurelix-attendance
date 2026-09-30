from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import secrets
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Header, HTTPException, Query
from fastapi.responses import Response
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from app.core.config import get_settings
from app.core.security import decode_access_token, require_admin
from app.db.mongodb import get_db
from app.services.photo_service import PhotoExpiredError, cleanup_expired_photos, open_photo

optional_bearer = HTTPBearer(auto_error=False)


def require_photo_cleanup_access(
    credentials: HTTPAuthorizationCredentials | None = Depends(optional_bearer),
    x_photo_cleanup_secret: str | None = Header(default=None),
) -> dict:
    settings = get_settings()
    configured_secret = settings.photo_cleanup_secret
    if configured_secret:
        if x_photo_cleanup_secret and secrets.compare_digest(x_photo_cleanup_secret, configured_secret):
            return {"role": "cron"}
        if credentials and secrets.compare_digest(credentials.credentials, configured_secret):
            return {"role": "cron"}
    if credentials:
        try:
            claims = decode_access_token(credentials.credentials)
        except HTTPException:
            claims = None
        if claims and claims.get("role") == "admin":
            return claims
    raise HTTPException(status_code=403, detail="Photo cleanup requires administrator access or a valid cleanup secret")

router = APIRouter(prefix="/api/admin", tags=["admin"])
KOLKATA_TIME_ZONE = ZoneInfo("Asia/Kolkata")

def format_location(location: dict | None) -> str:
    if not location or location.get("latitude") is None or location.get("longitude") is None:
        return ""
    if location.get("display_name"):
        return location["display_name"]
    if location.get("area"):
        city = location.get("city")
        return f"{location['area']}, {city}" if city and city not in location["area"] else location["area"]
    accuracy = location.get("accuracy")
    suffix = f", +/- {round(accuracy)} m" if accuracy is not None else ""
    return f"{location['latitude']:.5f}, {location['longitude']:.5f}{suffix}"


def excel_time_value(value):
    if value is None or value == "":
        return "—"
    try:
        timestamp = datetime.fromisoformat(value) if isinstance(value, str) else value
        if not isinstance(timestamp, datetime):
            return "—"
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=timezone.utc)
        formatted_time = timestamp.astimezone(KOLKATA_TIME_ZONE).strftime("%I:%M:%S %p").lower()
        return f"{formatted_time} IST"
    except (TypeError, ValueError, OverflowError):
        return "—"


def excel_working_hours(check_in, check_out):
    if not check_in or not check_out:
        return "—"
    try:
        check_in_time = datetime.fromisoformat(check_in) if isinstance(check_in, str) else check_in
        check_out_time = datetime.fromisoformat(check_out) if isinstance(check_out, str) else check_out
        duration = check_out_time - check_in_time
        return duration if duration >= timedelta(0) else "—"
    except (TypeError, ValueError, OverflowError):
        return "—"


def resolve_export_date_range(
    export_range: Literal["day", "custom", "month", "months", "year"],
    anchor_date: date | None,
    start_date: date | None,
    end_date: date | None,
) -> tuple[date, date]:
    """Normalize export inputs into the server-controlled inclusive date range."""
    if export_range == "custom":
        if not start_date or not end_date:
            raise HTTPException(status_code=422, detail="Start and end dates are required for a custom date range")
        if start_date > end_date:
            raise HTTPException(status_code=422, detail="Start date must not be after end date")
        return start_date, end_date

    if export_range == "months":
        if not start_date or not end_date:
            raise HTTPException(status_code=422, detail="Start and end months are required for a multiple-month export")
        start_month = start_date.replace(day=1)
        end_month_start = end_date.replace(day=1)
        if start_month > end_month_start:
            raise HTTPException(status_code=422, detail="Start month must not be after end month")
        next_month = (end_month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
        return start_month, next_month - timedelta(days=1)

    selected_date = anchor_date or date.today()
    if export_range == "day":
        return selected_date, selected_date
    if export_range == "month":
        start_date = selected_date.replace(day=1)
        next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
        return start_date, next_month - timedelta(days=1)
    return selected_date.replace(month=1, day=1), selected_date.replace(month=12, day=31)


def export_filename(export_range: str, start_date: date, end_date: date) -> str:
    if export_range == "day":
        suffix = start_date.isoformat()
    elif export_range == "custom":
        suffix = f"{start_date.isoformat()}-to-{end_date.isoformat()}"
    elif export_range == "month":
        suffix = start_date.strftime("%Y-%m")
    elif export_range == "months":
        suffix = f"{start_date.strftime('%Y-%m')}-to-{end_date.strftime('%Y-%m')}"
    else:
        suffix = str(start_date.year)
    return f"aurelix-attendance-{export_range}-{suffix}.xlsx"


def embed_attendance_photo(sheet, cell_coordinate: str, reference: dict | None) -> bool:
    """Embed a retained GridFS photo without exposing its storage identifier."""
    cell = sheet[cell_coordinate]
    if not reference:
        cell.value = "Expired / unavailable"
        return False
    try:
        stream = open_photo(reference)
        try:
            photo_data = stream.read()
        finally:
            close = getattr(stream, "close", None)
            if close:
                close()
        with Image.open(BytesIO(photo_data)) as source:
            thumbnail = source.convert("RGB")
            thumbnail.thumbnail((200, 150), Image.Resampling.LANCZOS)
            buffer = BytesIO()
            thumbnail.save(buffer, format="PNG", optimize=True)
        buffer.seek(0)
        image = ExcelImage(buffer)
        image.anchor = cell_coordinate
        sheet.add_image(image)
        return True
    except (PhotoExpiredError, OSError, ValueError, KeyError, TypeError):
        cell.value = "Expired / unavailable"
        return False
    except Exception:
        # Exports must remain available when a GridFS file was deleted or corrupted.
        cell.value = "Expired / unavailable"
        return False

@router.get("/export")
def export_attendance(
    export_range: Literal["day", "custom", "month", "months", "year"] = Query("day", alias="range"),
    anchor_date: date | None = Query(default=None, alias="date"),
    start_date: date | None = Query(default=None),
    end_date: date | None = Query(default=None),
    _claims: dict = Depends(require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    start_date, end_date = resolve_export_date_range(export_range, anchor_date, start_date, end_date)

    db = get_db()
    records = list(db.attendance.find(
        {"date": {"$gte": start_date.isoformat(), "$lte": end_date.isoformat()}},
        {"_id": 0},
    ).sort([("date", -1), ("employee_id", 1)]))
    active_employees = list(db.employees.find(
        {"role": "employee", "is_active": True},
        {"employee_id": 1, "full_name": 1, "email": 1, "department": 1},
    ))
    active_employees.sort(key=lambda employee: employee.get("employee_id") or "")
    employee_ids = {record.get("employee_id") for record in records}
    employees = {
        employee["employee_id"]: employee
        for employee in db.employees.find(
            {"employee_id": {"$in": list(employee_ids)}},
            {"employee_id": 1, "full_name": 1, "email": 1, "department": 1},
        )
    }
    employees.update({employee["employee_id"]: employee for employee in active_employees})

    records_by_employee_date = {}
    for record in records:
        key = (record.get("employee_id"), record.get("date"))
        records_by_employee_date.setdefault(key, []).append(record)

    export_records = []
    for day_offset in range((end_date - start_date).days + 1):
        current_date = (start_date + timedelta(days=day_offset)).isoformat()
        for employee in active_employees:
            employee_id = employee["employee_id"]
            matching_records = records_by_employee_date.pop((employee_id, current_date), [])
            if matching_records:
                export_records.extend(matching_records)
            else:
                export_records.append({
                    "date": current_date,
                    "employee_id": employee_id,
                    "user_name": employee.get("full_name"),
                    "check_in_time": None,
                    "check_in_location": None,
                    "check_out_time": None,
                    "check_out_location": None,
                    "final_status": "ABSENT",
                })

    for remaining_records in records_by_employee_date.values():
        export_records.extend(remaining_records)
    export_records.sort(key=lambda record: str(record.get("employee_id") or ""))
    export_records.sort(key=lambda record: str(record.get("date") or ""), reverse=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    headers = ["Date", "Employee ID", "Employee", "Email", "Department", "Check In", "Check In Location", "Check-in Photo", "Check Out", "Working Hours", "Check Out Location", "Check-out Photo", "Status"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102B42")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for record in export_records:
        employee = employees.get(record.get("employee_id"), {})
        row_number = sheet.max_row + 1
        is_absent = record.get("final_status") == "ABSENT"
        working_hours = excel_working_hours(record.get("check_in_time"), record.get("check_out_time"))
        sheet.append([
            record.get("date"), record.get("employee_id"), employee.get("full_name") or record.get("user_name") or record.get("employee_id"), employee.get("email", ""),
            employee.get("department", ""), excel_time_value(record.get("check_in_time")), "—" if is_absent and not record.get("check_in_location") else format_location(record.get("check_in_location")),
            None,
            excel_time_value(record.get("check_out_time")), working_hours, "—" if is_absent and not record.get("check_out_location") else format_location(record.get("check_out_location")),
            None,
            record.get("final_status"),
        ])
        sheet[f"J{row_number}"].number_format = "[h]:mm:ss"
        has_check_in_photo = embed_attendance_photo(sheet, f"H{row_number}", record.get("check_in_photo_reference"))
        has_check_out_photo = embed_attendance_photo(sheet, f"L{row_number}", record.get("check_out_photo_reference"))
        if has_check_in_photo or has_check_out_photo:
            sheet.row_dimensions[row_number].height = 115
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [14, 16, 24, 30, 20, 23, 28, 31, 23, 14, 28, 31, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    filename = export_filename(export_range, start_date, end_date)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.api_route("/photos/cleanup", methods=["GET", "POST"])
def run_expired_photo_cleanup(_access: dict = Depends(require_photo_cleanup_access)):
    result = cleanup_expired_photos()
    return {"message": "Expired attendance photos were cleaned up", **result}
